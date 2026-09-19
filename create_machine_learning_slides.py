from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTFILE = "machine_learning_presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(21, 34, 56)
TEAL = RGBColor(0, 150, 136)
GOLD = RGBColor(244, 180, 0)
RED = RGBColor(219, 68, 55)
LIGHT = RGBColor(246, 248, 251)
MUTED = RGBColor(95, 111, 130)
WHITE = RGBColor(255, 255, 255)


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(8.9), Inches(0.75))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.68), Inches(1.03), Inches(8.4), Inches(0.45))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = MUTED


def add_footer(slide, number):
    box = slide.shapes.add_textbox(Inches(11.85), Inches(7.0), Inches(0.75), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"{number:02}"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, items, x, y, w, h, font_size=21, color=NAVY):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.line_spacing = 1.1
    return box


def add_card(slide, x, y, w, h, title, body, accent=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(224, 230, 238)
    shape.line.width = Pt(1)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.12), Inches(h)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(x + 0.28), Inches(y + 0.22), Inches(w - 0.45), Inches(0.35))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(18)
    tp.font.bold = True
    tp.font.color.rgb = NAVY

    body_box = slide.shapes.add_textbox(Inches(x + 0.28), Inches(y + 0.72), Inches(w - 0.45), Inches(h - 0.85))
    bp = body_box.text_frame.paragraphs[0]
    bp.text = body
    bp.font.size = Pt(13.5)
    bp.font.color.rgb = MUTED
    bp.line_spacing = 1.15


def add_circle_label(slide, x, y, size, text, color):
    c = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size)
    )
    c.fill.solid()
    c.fill.fore_color.rgb = color
    c.line.fill.background()
    tf = c.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = 3


def add_arrow(slide, x, y, w, h, color=MUTED):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
title = slide.shapes.add_textbox(Inches(0.75), Inches(1.35), Inches(7.2), Inches(1.0))
p = title.text_frame.paragraphs[0]
p.text = "Machine Learning"
p.font.size = Pt(52)
p.font.bold = True
p.font.color.rgb = WHITE
sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.35), Inches(6.7), Inches(0.7))
sp = sub.text_frame.paragraphs[0]
sp.text = "Teaching computers to learn patterns from data"
sp.font.size = Pt(24)
sp.font.color.rgb = RGBColor(207, 226, 238)
add_circle_label(slide, 8.5, 1.3, 1.4, "DATA", TEAL)
add_circle_label(slide, 9.9, 2.75, 1.4, "MODEL", GOLD)
add_circle_label(slide, 8.5, 4.2, 1.4, "PREDICT", RED)
add_arrow(slide, 9.35, 2.0, 1.0, 0.35, RGBColor(119, 203, 196))
add_arrow(slide, 9.35, 4.05, 1.0, 0.35, RGBColor(239, 158, 151))

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "What Is Machine Learning?", "A simple definition")
add_bullets(
    slide,
    [
        "Machine learning is a branch of artificial intelligence.",
        "It allows computers to learn from examples instead of only following fixed rules.",
        "A trained model can make predictions, classify information, or discover patterns.",
    ],
    0.9,
    1.8,
    7.1,
    3.8,
)
add_card(slide, 8.55, 1.75, 3.8, 2.8, "Simple formula", "Data + algorithm = trained model\n\nTrained model + new input = prediction", TEAL)
add_footer(slide, 2)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "How Machine Learning Works", "The basic flow from data to prediction")
labels = [("Collect\nData", TEAL), ("Train\nModel", GOLD), ("Test\nAccuracy", RED), ("Use for\nPrediction", NAVY)]
x_positions = [0.95, 3.75, 6.55, 9.35]
for idx, (label, color) in enumerate(labels):
    add_circle_label(slide, x_positions[idx], 2.55, 1.55, label, color)
    if idx < 3:
        add_arrow(slide, x_positions[idx] + 1.7, 3.08, 0.85, 0.28)
add_bullets(
    slide,
    [
        "The model studies training data to learn useful patterns.",
        "Testing checks whether it works well on new examples.",
        "Good machine learning depends on good, representative data.",
    ],
    1.15,
    5.05,
    10.5,
    1.4,
    17,
)
add_footer(slide, 3)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "Main Types of Machine Learning")
add_card(slide, 0.8, 1.7, 3.75, 3.5, "Supervised learning", "Learns from labeled examples.\n\nExample: predicting house prices from previous house sales.", TEAL)
add_card(slide, 4.8, 1.7, 3.75, 3.5, "Unsupervised learning", "Finds hidden patterns without answer labels.\n\nExample: grouping customers by buying behavior.", GOLD)
add_card(slide, 8.8, 1.7, 3.75, 3.5, "Reinforcement learning", "Learns by receiving rewards or penalties.\n\nExample: an AI learning to play a game.", RED)
add_footer(slide, 4)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "Everyday Examples")
items = [
    ("Recommendations", "Netflix, YouTube, TikTok, and shopping apps suggest content or products."),
    ("Voice assistants", "Speech recognition helps systems understand spoken commands."),
    ("Fraud detection", "Banks detect unusual transactions and warn customers."),
    ("Medical support", "Models can help identify patterns in scans and patient records."),
]
for i, (title_text, body) in enumerate(items):
    row = i // 2
    col = i % 2
    add_card(slide, 0.9 + col * 6.0, 1.55 + row * 2.3, 5.25, 1.75, title_text, body, [TEAL, GOLD, RED, NAVY][i])
add_footer(slide, 5)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "Why Machine Learning Is Useful")
add_bullets(
    slide,
    [
        "Handles large amounts of data faster than humans can.",
        "Finds patterns that may be difficult to see manually.",
        "Improves decisions in business, science, health, and security.",
        "Can adapt when new training data becomes available.",
    ],
    0.95,
    1.7,
    6.5,
    4.4,
)
add_card(slide, 8.0, 1.65, 4.2, 3.5, "Key idea", "Machine learning is most valuable when there is enough quality data and a clear problem to solve.", TEAL)
add_footer(slide, 6)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title(slide, "Challenges and Risks")
add_card(slide, 0.9, 1.65, 3.65, 3.7, "Bad data", "If the data is incomplete, wrong, or biased, the model may learn poor patterns.", RED)
add_card(slide, 4.85, 1.65, 3.65, 3.7, "Overfitting", "A model may memorize training examples and perform badly on new data.", GOLD)
add_card(slide, 8.8, 1.65, 3.65, 3.7, "Privacy and fairness", "Systems must protect personal data and avoid unfair decisions.", TEAL)
add_footer(slide, 7)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
box = slide.shapes.add_textbox(Inches(0.85), Inches(0.85), Inches(7.7), Inches(0.8))
p = box.text_frame.paragraphs[0]
p.text = "Summary"
p.font.size = Pt(42)
p.font.bold = True
p.font.color.rgb = WHITE
add_bullets(
    slide,
    [
        "Machine learning teaches computers to learn from data.",
        "It is used for predictions, classification, recommendations, and pattern discovery.",
        "The three major types are supervised, unsupervised, and reinforcement learning.",
        "Good results require good data, careful testing, and responsible use.",
    ],
    0.95,
    2.0,
    7.8,
    3.8,
    20,
    WHITE,
)
add_circle_label(slide, 9.4, 2.1, 1.45, "LEARN", TEAL)
add_circle_label(slide, 10.45, 3.35, 1.45, "TEST", GOLD)
add_circle_label(slide, 9.4, 4.6, 1.45, "USE", RED)

prs.save(OUTFILE)
print(OUTFILE)
